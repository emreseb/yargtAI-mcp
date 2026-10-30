"""Dilekçeyi biçimli bir Word (.docx) belgesine döker.

Düz metin yerine avukatın Word'de açıp düzenleyebileceği gerçek bir belge:
Times New Roman 12, ortalı mahkeme başlığı, hizalı taraf bloğu, kalın bölüm
başlıkları, iki yana yaslı gövde, sağa yaslı imza. Boş bırakılan zorunlu alanlar
sarı vurgulu "[ … buraya yazınız ]" yer tutucusu olarak gelir — doldurulacak slot.
"""

from __future__ import annotations

from datetime import datetime
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Cm, Pt

from . import dilekce
from .config import DILEKCE_DIR

_LABEL_STOP = Cm(4.0)


def _doc():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    return doc


def _add_val(p, c: dict[str, Any]) -> None:
    """Çözülmüş alanı paragrafa yazar; yer tutucuysa sarı vurgulu."""
    if c["placeholder"]:
        r = p.add_run(f"[ {c['aciklama']} — buraya yazınız ]")
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    else:
        for i, satir in enumerate(str(c["deger"]).split("\n")):
            if i:
                p.add_run().add_break()
            p.add_run(satir)


def _heading(doc, c_mahkeme) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    if c_mahkeme["placeholder"]:
        r = p.add_run(f"[ {c_mahkeme['aciklama']} — buraya yazınız ]")
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    else:
        r = p.add_run(str(c_mahkeme["deger"]) + "'NE")
    r.bold = True


def _row(doc, label: str, parts: list) -> None:
    """Etiket : değer satırı (hizalı, asılı girinti). parts: cozum, metin ya da
    (önek, cozum) öğelerinden oluşan liste; her öğe yeni satıra geçer."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = _LABEL_STOP
    pf.first_line_indent = -_LABEL_STOP
    pf.space_after = Pt(2)
    pf.tab_stops.add_tab_stop(_LABEL_STOP)
    p.add_run(label).bold = True
    p.add_run("\t: ")
    for i, part in enumerate(parts):
        if i:
            p.add_run().add_break()
        if isinstance(part, tuple):
            onek, c = part
            if onek:
                p.add_run(onek)
            _add_val(p, c)
        elif isinstance(part, dict):
            _add_val(p, part)
        else:
            p.add_run(str(part))


def _block(doc, baslik: str, c) -> None:
    """Başlık + iki yana yaslı gövde (uzun bölümler: açıklamalar, sonuç/istem)."""
    hp = doc.add_paragraph()
    hp.paragraph_format.space_before = Pt(10)
    hp.add_run(baslik + " :").bold = True
    bp = doc.add_paragraph()
    bp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    bp.paragraph_format.space_after = Pt(6)
    _add_val(bp, c)


def _signature(doc, c_tarih, rol: str, c_imza) -> None:
    doc.add_paragraph()
    for content in (c_tarih, rol, c_imza):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(0)
        if isinstance(content, str):
            p.add_run(content).bold = True
        else:
            _add_val(p, content)


def _build_dava(doc, c) -> None:
    _heading(doc, c["mahkeme"])
    _row(doc, "DAVACI", [c["davaci_ad"], ("T.C. Kimlik No: ", c["davaci_tc"]), c["davaci_adres"]])
    _row(doc, "VEKİLİ", [c["vekil"]])
    _row(doc, "DAVALI", [c["davali_ad"], c["davali_adres"]])
    _row(doc, "KONU", [c["konu"]])
    _row(doc, "DAVA DEĞERİ", [c["dava_degeri"]])
    _block(doc, "AÇIKLAMALAR", c["aciklamalar"])
    _row(doc, "HUKUKİ SEBEPLER", [c["hukuki_sebepler"]])
    _row(doc, "DELİLLER", [c["deliller"]])
    _block(doc, "SONUÇ VE İSTEM", c["talep"])
    _signature(doc, c["tarih"], "Davacı / Vekili", c["vekil"])


def _build_cevap(doc, c) -> None:
    _heading(doc, c["mahkeme"])
    _row(doc, "DOSYA NO", [c["dosya_no"]])
    _row(doc, "CEVAP VEREN\n(DAVALI)", [c["davali_ad"], ("T.C. Kimlik No: ", c["davali_tc"]), c["davali_adres"]])
    _row(doc, "VEKİLİ", [c["vekil"]])
    _row(doc, "DAVACI", [c["davaci_ad"], c["davaci_adres"]])
    _row(doc, "KONU", [c["konu"]])
    _block(doc, "AÇIKLAMALAR (SAVUNMA)", c["aciklamalar"])
    _row(doc, "HUKUKİ SEBEPLER", [c["hukuki_sebepler"]])
    _row(doc, "DELİLLER", [c["deliller"]])
    _block(doc, "SONUÇ VE İSTEM", c["talep"])
    _signature(doc, c["tarih"], "Davalı / Vekili", c["vekil"])


_BUILDERS = {"dava_dilekcesi": _build_dava, "cevap_dilekcesi": _build_cevap}


def belge_olustur(
    tur: str, alanlar: dict[str, Any] | None = None, dosya_adi: str | None = None
) -> dict[str, Any]:
    """Şablonu .docx olarak üretir, diske yazar, dosya yolunu döndürür."""
    s, cozum, eksik = dilekce.coz(tur, alanlar)
    builder = _BUILDERS.get(tur)
    if not builder:
        raise ValueError(f"{tur!r} için .docx üretici yok. Seçenekler: {', '.join(_BUILDERS)}")

    doc = _doc()
    builder(doc, cozum)

    DILEKCE_DIR.mkdir(parents=True, exist_ok=True)
    if not dosya_adi:
        dosya_adi = f"{tur}_{datetime.now():%Y%m%d_%H%M%S}.docx"
    if not dosya_adi.lower().endswith(".docx"):
        dosya_adi += ".docx"
    yol = (DILEKCE_DIR / dosya_adi).resolve()
    doc.save(str(yol))

    return {
        "tur": tur,
        "ad": s["ad"],
        "dosya_yolu": str(yol),
        "eksik_alanlar": eksik,
        "not": (
            "Word (.docx) taslağı oluşturuldu. Sarı vurgulu "
            "'[ … buraya yazınız ]' alanları doldurulmayı bekler. "
            "Göndermeden önce avukat kontrolünden geçirin."
        ),
    }
